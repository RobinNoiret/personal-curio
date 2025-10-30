"""
Script pour inspecter le template et voir tous les paragraphes
"""
from docx import Document

template_path = r"c:\Users\Noiret\Documents\0 - Tools\prosit_manager\pbl-cer_template_2025.docx"

doc = Document(template_path)

print("=== Inspection du template ===\n")
print(f"Nombre total de paragraphes: {len(doc.paragraphs)}\n")

# Chercher autour de "Hypothèses" et "Plan d'actions"
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if any(keyword in text.lower() for keyword in ['hypothèse', 'plan', 'recherche']):
        print(f"[{i}] Style: {p.style.name:30} | Texte: {text}")
